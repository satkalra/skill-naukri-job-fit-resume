#!/usr/bin/env python3
"""
build_resume.py — turn a resume JSON (see resume_schema.json) into an
ATS-safe DOCX, and optionally a PDF.

Usage:
    python3 build_resume.py input.json --out-dir DIR [--pdf]

Design rules enforced (see references/ats-guidelines.md §2-3):
    - single column, left aligned
    - standard section headings
    - Month YYYY dates, validated and rejected if mixed/malformed
    - contact info in the document body (not header/footer)
    - standard round bullets (python-docx "List Bullet" style)
    - no tables, text boxes, or graphics
    - standard font (Calibri), 14pt name / 10.5-11pt body
    - 0.6" margins
    - output filename: First_Last_TargetRole_Resume.{docx,pdf}
"""

import argparse
import json
import re
import sys
from pathlib import Path

MONTH_RE = re.compile(
    r"^(January|February|March|April|May|June|July|August|September|October|November|December) \d{4}$"
)


def fail(msg):
    print(f"ERROR: {msg}", file=sys.stderr)
    sys.exit(1)


def validate_date(value, field_name):
    if value == "Present":
        return
    if not MONTH_RE.match(value):
        fail(
            f"Invalid date '{value}' for {field_name}. "
            "Dates must be 'Month YYYY' (e.g. 'March 2022') or 'Present'. "
            "Mixed/malformed date formats break ATS parsing (see §2)."
        )


def validate_data(data):
    required_top = ["name", "contact", "target_role", "experience", "skills", "education"]
    for key in required_top:
        if key not in data:
            fail(f"Missing required field '{key}' in input JSON.")

    contact = data["contact"]
    for key in ["email", "phone", "location"]:
        if key not in contact:
            fail(f"Missing required contact field '{key}'.")

    if not data["experience"]:
        fail("At least one experience entry is required.")

    for i, exp in enumerate(data["experience"]):
        for key in ["title", "employer", "start_date", "end_date", "bullets"]:
            if key not in exp:
                fail(f"Experience entry {i} missing required field '{key}'.")
        validate_date(exp["start_date"], f"experience[{i}].start_date")
        validate_date(exp["end_date"], f"experience[{i}].end_date")

    for i, edu in enumerate(data["education"]):
        for key in ["degree", "institution", "date"]:
            if key not in edu:
                fail(f"Education entry {i} missing required field '{key}'.")
        validate_date(edu["date"], f"education[{i}].date")


def safe_filename_component(s):
    return re.sub(r"[^A-Za-z0-9]+", "_", s.strip()).strip("_")


def build_filename(data):
    name_parts = data["name"].strip().split()
    first = name_parts[0] if name_parts else "Candidate"
    last = name_parts[-1] if len(name_parts) > 1 else ""
    role = safe_filename_component(data["target_role"])
    base = "_".join(p for p in [safe_filename_component(first), safe_filename_component(last), role, "Resume"] if p)
    return base


def build_docx(data, out_path):
    try:
        from docx import Document
        from docx.shared import Pt, Inches
        from docx.enum.text import WD_ALIGN_PARAGRAPH
    except ImportError:
        fail(
            "python-docx is not installed. Install it with 'pip install python-docx', "
            "or use the no-code fallback (plain-text ATS-safe resume)."
        )

    doc = Document()

    # 0.6" margins, single column (default docx section is already single column)
    for section in doc.sections:
        section.left_margin = Inches(0.6)
        section.right_margin = Inches(0.6)
        section.top_margin = Inches(0.6)
        section.bottom_margin = Inches(0.6)

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)

    # Name
    name_p = doc.add_paragraph()
    name_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = name_p.add_run(data["name"])
    run.bold = True
    run.font.size = Pt(14)
    run.font.name = "Calibri"

    # Contact line — in the body, never header/footer
    contact = data["contact"]
    contact_bits = [contact["location"], contact["phone"], contact["email"]]
    if contact.get("linkedin"):
        contact_bits.append(contact["linkedin"])
    if contact.get("portfolio"):
        contact_bits.append(contact["portfolio"])
    contact_p = doc.add_paragraph(" | ".join(contact_bits))
    contact_p.paragraph_format.space_after = Pt(6)

    # Headline
    if data.get("headline"):
        h = doc.add_paragraph()
        r = h.add_run(data["headline"])
        r.bold = True
        r.font.size = Pt(11)
        h.paragraph_format.space_after = Pt(6)

    def add_heading(text):
        p = doc.add_paragraph()
        r = p.add_run(text.upper())
        r.bold = True
        r.font.size = Pt(12)
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(4)

    # Summary
    if data.get("summary"):
        add_heading("Professional Summary")
        doc.add_paragraph(data["summary"])

    # Skills
    add_heading("Skills")
    for group in data["skills"]:
        p = doc.add_paragraph()
        r = p.add_run(f"{group['category']}: ")
        r.bold = True
        p.add_run(", ".join(group["items"]))

    # Experience
    add_heading("Professional Experience")
    for exp in data["experience"]:
        p = doc.add_paragraph()
        r = p.add_run(exp["title"])
        r.bold = True
        loc = f" — {exp['location']}" if exp.get("location") else ""
        p.add_run(f", {exp['employer']}{loc}")
        date_p = doc.add_paragraph(f"{exp['start_date']} – {exp['end_date']}")
        date_p.paragraph_format.space_after = Pt(2)
        for bullet in exp["bullets"]:
            bp = doc.add_paragraph(bullet, style="List Bullet")
            bp.paragraph_format.space_after = Pt(2)

    # Projects
    if data.get("projects"):
        add_heading("Projects")
        for proj in data["projects"]:
            p = doc.add_paragraph()
            r = p.add_run(proj["name"])
            r.bold = True
            if proj.get("link"):
                p.add_run(f" ({proj['link']})")
            for bullet in proj["bullets"]:
                bp = doc.add_paragraph(bullet, style="List Bullet")
                bp.paragraph_format.space_after = Pt(2)

    # Education
    add_heading("Education")
    for edu in data["education"]:
        p = doc.add_paragraph()
        r = p.add_run(edu["degree"])
        r.bold = True
        loc = f" — {edu['location']}" if edu.get("location") else ""
        p.add_run(f", {edu['institution']}{loc}")
        doc.add_paragraph(edu["date"])

    # Certifications
    if data.get("certifications"):
        add_heading("Certifications")
        for cert in data["certifications"]:
            line = f"{cert['name']}, {cert['issuer']}"
            if cert.get("date"):
                line += f" ({cert['date']})"
            doc.add_paragraph(line)

    doc.save(out_path)


def build_pdf_via_soffice(docx_path, out_dir):
    import shutil
    import subprocess

    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        return False
    try:
        subprocess.run(
            [soffice, "--headless", "--convert-to", "pdf", "--outdir", str(out_dir), str(docx_path)],
            check=True,
            capture_output=True,
            timeout=120,
        )
        return True
    except Exception as e:
        print(f"WARNING: soffice PDF conversion failed ({e}); falling back to reportlab.", file=sys.stderr)
        return False


def build_pdf_via_reportlab(data, out_path):
    try:
        from reportlab.lib.pagesizes import LETTER
        from reportlab.lib.units import inch
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, ListFlowable, ListItem
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib.enums import TA_LEFT
    except ImportError:
        fail(
            "Neither LibreOffice (soffice) nor reportlab is available to produce a PDF. "
            "Install reportlab with 'pip install reportlab', or omit --pdf and hand-export the DOCX to PDF."
        )

    styles = getSampleStyleSheet()
    name_style = ParagraphStyle("NameStyle", parent=styles["Normal"], fontName="Helvetica-Bold", fontSize=14, alignment=TA_LEFT)
    heading_style = ParagraphStyle("HeadingStyle", parent=styles["Normal"], fontName="Helvetica-Bold", fontSize=12, spaceBefore=10, spaceAfter=4, alignment=TA_LEFT)
    body_style = ParagraphStyle("BodyStyle", parent=styles["Normal"], fontName="Helvetica", fontSize=10.5, alignment=TA_LEFT)
    bold_body_style = ParagraphStyle("BoldBodyStyle", parent=body_style, fontName="Helvetica-Bold")

    doc = SimpleDocTemplate(
        str(out_path), pagesize=LETTER,
        leftMargin=0.6 * inch, rightMargin=0.6 * inch,
        topMargin=0.6 * inch, bottomMargin=0.6 * inch,
    )
    story = [Paragraph(data["name"], name_style)]

    contact = data["contact"]
    contact_bits = [contact["location"], contact["phone"], contact["email"]]
    if contact.get("linkedin"):
        contact_bits.append(contact["linkedin"])
    if contact.get("portfolio"):
        contact_bits.append(contact["portfolio"])
    story.append(Paragraph(" | ".join(contact_bits), body_style))
    story.append(Spacer(1, 6))

    if data.get("headline"):
        story.append(Paragraph(data["headline"], bold_body_style))
        story.append(Spacer(1, 6))

    def heading(text):
        story.append(Paragraph(text.upper(), heading_style))

    if data.get("summary"):
        heading("Professional Summary")
        story.append(Paragraph(data["summary"], body_style))

    heading("Skills")
    for group in data["skills"]:
        story.append(Paragraph(f"<b>{group['category']}:</b> {', '.join(group['items'])}", body_style))

    heading("Professional Experience")
    for exp in data["experience"]:
        loc = f" — {exp['location']}" if exp.get("location") else ""
        story.append(Paragraph(f"<b>{exp['title']}</b>, {exp['employer']}{loc}", body_style))
        story.append(Paragraph(f"{exp['start_date']} – {exp['end_date']}", body_style))
        items = [ListItem(Paragraph(b, body_style)) for b in exp["bullets"]]
        story.append(ListFlowable(items, bulletType="bullet"))

    if data.get("projects"):
        heading("Projects")
        for proj in data["projects"]:
            title = proj["name"] + (f" ({proj['link']})" if proj.get("link") else "")
            story.append(Paragraph(f"<b>{title}</b>", body_style))
            items = [ListItem(Paragraph(b, body_style)) for b in proj["bullets"]]
            story.append(ListFlowable(items, bulletType="bullet"))

    heading("Education")
    for edu in data["education"]:
        loc = f" — {edu['location']}" if edu.get("location") else ""
        story.append(Paragraph(f"<b>{edu['degree']}</b>, {edu['institution']}{loc}", body_style))
        story.append(Paragraph(edu["date"], body_style))

    if data.get("certifications"):
        heading("Certifications")
        for cert in data["certifications"]:
            line = f"{cert['name']}, {cert['issuer']}"
            if cert.get("date"):
                line += f" ({cert['date']})"
            story.append(Paragraph(line, body_style))

    doc.build(story)


def main():
    parser = argparse.ArgumentParser(description="Build an ATS-safe resume DOCX (and optionally PDF) from JSON.")
    parser.add_argument("input", help="Path to input resume JSON (see resume_schema.json)")
    parser.add_argument("--out-dir", required=True, help="Directory to write output files into")
    parser.add_argument("--pdf", action="store_true", help="Also produce a PDF")
    args = parser.parse_args()

    input_path = Path(args.input)
    if not input_path.exists():
        fail(f"Input file not found: {input_path}")

    try:
        data = json.loads(input_path.read_text())
    except json.JSONDecodeError as e:
        fail(f"Input file is not valid JSON: {e}")

    validate_data(data)

    out_dir = Path(args.out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)

    base = build_filename(data)
    docx_path = out_dir / f"{base}.docx"
    build_docx(data, docx_path)
    print(f"Wrote {docx_path}")

    if args.pdf:
        pdf_path = out_dir / f"{base}.pdf"
        if build_pdf_via_soffice(docx_path, out_dir):
            # soffice names the output after the docx stem
            produced = out_dir / f"{docx_path.stem}.pdf"
            if produced != pdf_path and produced.exists():
                produced.rename(pdf_path)
            print(f"Wrote {pdf_path} (via LibreOffice)")
        else:
            build_pdf_via_reportlab(data, pdf_path)
            print(f"Wrote {pdf_path} (via reportlab fallback)")


if __name__ == "__main__":
    main()
